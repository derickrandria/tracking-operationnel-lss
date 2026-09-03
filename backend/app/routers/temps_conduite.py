# -*- coding: utf-8 -*-
"""Module « Temps de conduite » — Compteur TCH et matrice journalière par chauffeur.

Règles métier :
  1. Le TCH (Temps de Conduite Hebdomadaire) = somme des TCJ du chauffeur
     depuis son dernier reset TCH.
  2. Limite maximale = 56 heures (201 600 s) ; TCH restant = 56h - TCH cumulé.
  3. RESET DU TCH : le compteur TCH est remis à zéro lorsqu'un chauffeur
     n'a effectué aucun trajet valide (≥ 0,3 km) pendant une période continue
     de 24 heures ou plus (86 400 s), tous véhicules confondus.
  4. TEMPS RÉEL : le TCJ d'aujourd'hui est dynamique et s'additionne au TCH
     des journées précédentes depuis le dernier reset.
  5. HISTORIQUE & SUIVI : réutilisation stricte des valeurs TCJ/TTJ existantes
     sans duplication de calcul.
  6. ALERTES :
     - TCH cumulé ≥ 46h00 (165 600 s) : « TCH proche de la limite — [Nom] — [Cumul] — [Restant] »
     - TCH cumulé ≥ 56h00 (201 600 s) : « TCH limite atteinte — [Nom] »
"""
import io
import re
from datetime import date, datetime, timedelta

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session, joinedload, selectinload

from ..config import calculer_tokens_set, jour_attribution, now_local
from ..database import get_db
from ..engine import get_seuils
from ..models import (Conducteur, HistoriqueJournalier, StatutConducteur,
                     StatutValidationTrajet, SuiviJournalier, Trajet, Vehicule)
from ..reparation import normaliser_libelle
from ..security import TOUS, require_roles
from ..serializers import fmt_hms, iso

router = APIRouter(prefix="/api/temps-conduite", tags=["temps-conduite"])

SEUIL_TCH_MAX_S = 201600    # 56h00
SEUIL_TCH_ALERTE_S = 165600 # 46h00
SEUIL_RESET_REPOS_S = 86400 # 24h00 de repos continu


def _dt_iso(texte) -> datetime | None:
    if not texte:
        return None
    try:
        return datetime.fromisoformat(str(texte))
    except (ValueError, TypeError):
        return None


def _determiner_periode(du: str | None, au: str | None) -> tuple[date, date]:
    """Période par défaut : les 7 derniers jours (J-6 -> aujourd'hui)."""
    fin = date.fromisoformat(au) if au else now_local().date()
    debut = date.fromisoformat(du) if du else fin - timedelta(days=6)
    if debut > fin:
        debut, fin = fin, debut
    return debut, fin


def extraire_donnees_chauffeurs(db: Session, debut_fenetre: date, fin_fenetre: date,
                                maintenant: datetime | None = None,
                                conducteur_id_filtre: str | None = None) -> dict:
    """Extrait et calcule l'ensemble des données TCH pour les chauffeurs.

    Règles :
      - Déduplication & réconciliation canonique (MZoneX / CamtrackPro / Fiches locales).
      - Détection des coupures de repos continu >= 24h.
      - Plafonnement au cycle hebdomadaire (max 6-7 jours glissants, jamais d'accumulation
        sur des semaines entières).
      - Affichage complet nom + prénom et réutilisation stricte des valeurs TCJ existantes.
    """
    maintenant = maintenant or now_local()
    jour_courant = jour_attribution(maintenant)
    borne_recul = min(debut_fenetre, fin_fenetre) - timedelta(days=21)

    # 1. Récupération de tous les conducteurs et construction de la table de déduplication canonique
    q_cond = select(Conducteur).order_by(Conducteur.nom_prenom)
    if conducteur_id_filtre:
        q_cond = q_cond.where(Conducteur.id == conducteur_id_filtre)
    conducteurs = db.scalars(q_cond).all()

    # Déduplication par nom normalisé : MZoneX vs CamtrackPro vs saisie manuelle
    groupes_canon: dict[str, list[Conducteur]] = {}
    id_vers_gardien: dict[str, Conducteur] = {}
    nom_norm_vers_gardien: dict[str, Conducteur] = {}
    tset_vers_gardien: dict[str, Conducteur] = {}

    for c in conducteurs:
        cle = (c.nom_normalise or normaliser_libelle(c.nom_prenom))[:170]
        groupes_canon.setdefault(cle, []).append(c)

    for cle, fiches in groupes_canon.items():
        # Sélection du gardien : matricule officiel CH... d'abord, puis téléphone, puis statut
        def score(f: Conducteur):
            s = 0
            if f.matricule and not f.matricule.startswith("AUTO-"):
                s += 500
            if f.telephone:
                s += 100
            if f.statut == StatutConducteur.ACTIF:
                s += 50
            return (-s, f.date_creation or datetime.min, f.id)

        gardien = sorted(fiches, key=score)[0]
        nom_norm_vers_gardien[cle] = gardien
        for f in fiches:
            id_vers_gardien[f.id] = gardien

    # Liste unique des fiches gardiennes actives pour la vue
    gardiens_uniques = list({g.id: g for g in id_vers_gardien.values()}.values())
    gardiens_uniques.sort(key=lambda x: (x.nom_prenom or "").upper())

    for g in gardiens_uniques:
        tset = getattr(g, "tokens_set", None) or (calculer_tokens_set(g.nom_prenom)[:170] if g.nom_prenom else "")
        if tset:
            tset_vers_gardien[tset] = g

    # 2. Récupération des SuiviJournalier (avec trajets et véhicule)
    q_suivi = (
        select(SuiviJournalier)
        .options(selectinload(SuiviJournalier.trajets), joinedload(SuiviJournalier.vehicule))
        .where(SuiviJournalier.date_jour >= borne_recul,
               SuiviJournalier.date_jour <= fin_fenetre)
        .order_by(SuiviJournalier.date_jour)
    )
    suivis = db.scalars(q_suivi).all()

    # 3. Récupération des HistoriqueJournalier
    q_hist = (
        select(HistoriqueJournalier)
        .options(joinedload(HistoriqueJournalier.vehicule))
        .where(HistoriqueJournalier.date_jour >= borne_recul,
               HistoriqueJournalier.date_jour <= fin_fenetre)
        .order_by(HistoriqueJournalier.date_jour)
    )
    historiques = db.scalars(q_hist).all()

    # Structure d'agrégation par chauffeur gardien :
    # gardien_id -> {
    #    "gardien": Conducteur,
    #    "jours": { date -> { "tcj_s": int, "ttj_s": int, "vehicules": set(), "en_cours": bool } },
    #    "spans": [ (debut_dt, fin_dt, plaque) ]
    # }
    data_chauffeurs: dict[str, dict] = {
        g.id: {"gardien": g, "jours": {}, "spans": []} for g in gardiens_uniques
    }

    def resoudre_gardien(c_id: str | None, nom_str: str | None) -> Conducteur | None:
        if c_id and c_id in id_vers_gardien:
            return id_vers_gardien[c_id]
        if nom_str:
            cle = normaliser_libelle(nom_str)[:170]
            if cle in nom_norm_vers_gardien:
                return nom_norm_vers_gardien[cle]
            tset = calculer_tokens_set(nom_str)[:170]
            if tset in tset_vers_gardien:
                return tset_vers_gardien[tset]
            mots_c = set(tset.split())
            if mots_c:
                for g in gardiens_uniques:
                    g_tokens = set((getattr(g, "tokens_set", None) or calculer_tokens_set(g.nom_prenom)).split())
                    if mots_c.issubset(g_tokens) or (g.prenom_usuel and normaliser_libelle(g.prenom_usuel) in mots_c):
                        return g
        return None

    # A. Intégration EXCLUSIVE de HistoriqueJournalier pour tous les jours passés archivés (j < jour_courant)
    jours_archives_vus = set()
    for h in historiques:
        j = h.date_jour
        if j >= jour_courant:
            continue  # Les jours courants ou futurs sont traités exclusivement via SuiviJournalier

        d = h.donnees or {}
        plaque = (h.vehicule.plaque if h.vehicule else None) or d.get("plaque")
        trajets_snap = d.get("trajets") or []
        tcj_archive = int(d.get("tcj_s") or 0)
        ttj_archive = int(d.get("ttj_s") or 0)

        trips_par_gardien: dict[str, list[dict]] = {}
        if trajets_snap:
            for t in trajets_snap:
                if not isinstance(t, dict):
                    continue
                if t.get("statut_validation") == "REJETE":
                    continue
                dist = t.get("distance_km")
                if dist is not None and dist < 0.3:
                    continue
                deb = _dt_iso(t.get("heure_debut"))
                if deb is None:
                    continue
                fin = _dt_iso(t.get("heure_fin"))
                duree = int((fin - deb).total_seconds()) if (fin and fin >= deb) else 0

                # Chauffeur spécifique du trajet :
                # 1. Badge explicite sur le trajet
                t_gard = resoudre_gardien(t.get("conducteur_badge_id"), t.get("conducteur_badge"))
                # 2. Si pas de badge dans le snapshot, chercher dans la table trajets DB
                if not t_gard and t.get("id"):
                    t_db = db.get(Trajet, str(t.get("id")))
                    if t_db:
                        t_gard = resoudre_gardien(t_db.conducteur_badge_id, t_db.conducteur_badge)
                # 3. Si pas de badge sur le trajet et que le camion a un titulaire officiel
                if not t_gard and h.vehicule and h.vehicule.conducteur_actuel_id:
                    t_gard = resoudre_gardien(h.vehicule.conducteur_actuel_id, None)
                # 4. Fallback sur le chauffeur de l'archive
                if not t_gard:
                    t_gard = resoudre_gardien(h.conducteur_id, d.get("chauffeur") or plaque)
                if not t_gard:
                    continue

                trips_par_gardien.setdefault(t_gard.id, []).append({
                    "gardien": t_gard,
                    "debut": deb,
                    "fin": fin or (deb + timedelta(seconds=max(300, duree))),
                    "duree": duree,
                    "plaque": plaque
                })

        if trips_par_gardien:
            nb_conducteurs = len(trips_par_gardien)
            if nb_conducteurs == 1:
                tgid, items = list(trips_par_gardien.items())[0]
                t_gard = items[0]["gardien"]
                if not conducteur_id_filtre or tgid == conducteur_id_filtre:
                    jours_archives_vus.add((tgid, j))
                    c_data = data_chauffeurs.setdefault(tgid, {"gardien": t_gard, "jours": {}, "spans": []})
                    for it in items:
                        c_data["spans"].append((it["debut"], it["fin"], plaque))
                    tcj_val = tcj_archive if tcj_archive > 0 else sum(it["duree"] for it in items)
                    ttj_val = max(ttj_archive, tcj_val)
                    if j not in c_data["jours"]:
                        c_data["jours"][j] = {
                            "tcj_s": tcj_val, "ttj_s": ttj_val,
                            "vehicules": {plaque} if plaque else set(),
                            "en_cours": False
                        }
                    else:
                        c_data["jours"][j]["tcj_s"] += tcj_val
                        c_data["jours"][j]["ttj_s"] = max(c_data["jours"][j]["ttj_s"] + ttj_val, c_data["jours"][j]["tcj_s"])
                        if plaque:
                            c_data["jours"][j]["vehicules"].add(plaque)
            else:
                for tgid, items in trips_par_gardien.items():
                    if conducteur_id_filtre and tgid != conducteur_id_filtre:
                        continue
                    t_gard = items[0]["gardien"]
                    jours_archives_vus.add((tgid, j))
                    c_data = data_chauffeurs.setdefault(tgid, {"gardien": t_gard, "jours": {}, "spans": []})
                    for it in items:
                        c_data["spans"].append((it["debut"], it["fin"], plaque))
                    tcj_c = sum(it["duree"] for it in items)
                    deb_min = min(it["debut"] for it in items)
                    fin_max = max(it["fin"] for it in items)
                    amp = int((fin_max - deb_min).total_seconds())
                    ttj_c = max(amp, tcj_c)
                    if j not in c_data["jours"]:
                        c_data["jours"][j] = {
                            "tcj_s": tcj_c, "ttj_s": ttj_c,
                            "vehicules": {plaque} if plaque else set(),
                            "en_cours": False
                        }
                    else:
                        c_data["jours"][j]["tcj_s"] += tcj_c
                        c_data["jours"][j]["ttj_s"] = max(c_data["jours"][j]["ttj_s"] + ttj_c, c_data["jours"][j]["tcj_s"])
                        if plaque:
                            c_data["jours"][j]["vehicules"].add(plaque)
        else:
            gardien = None
            if h.vehicule and h.vehicule.conducteur_actuel_id:
                gardien = resoudre_gardien(h.vehicule.conducteur_actuel_id, None)
            if not gardien:
                gardien = resoudre_gardien(h.conducteur_id, d.get("chauffeur") or plaque)
            if gardien:
                gid = gardien.id
                if not conducteur_id_filtre or gid == conducteur_id_filtre:
                    jours_archives_vus.add((gid, j))
                    c_data = data_chauffeurs.setdefault(gid, {"gardien": gardien, "jours": {}, "spans": []})
                    tcj_val = tcj_archive
                    ttj_val = max(ttj_archive, tcj_val)
                    if j not in c_data["jours"]:
                        c_data["jours"][j] = {
                            "tcj_s": tcj_val, "ttj_s": ttj_val,
                            "vehicules": {plaque} if plaque else set(),
                            "en_cours": False
                        }
                    else:
                        c_data["jours"][j]["tcj_s"] += tcj_val
                        c_data["jours"][j]["ttj_s"] = max(c_data["jours"][j]["ttj_s"] + ttj_val, c_data["jours"][j]["tcj_s"])
                        if plaque:
                            c_data["jours"][j]["vehicules"].add(plaque)

                    if tcj_val > 0:
                        h_dep = _dt_iso(d.get("heure_depart"))
                        if h_dep is not None:
                            fin_eff = h_dep + timedelta(seconds=max(1800, ttj_val or tcj_val))
                            c_data["spans"].append((h_dep, fin_eff, plaque))
                        else:
                            deb_def = datetime.combine(j, datetime.min.time()).replace(hour=6)
                            fin_def = deb_def + timedelta(seconds=max(1800, ttj_val or tcj_val))
                            c_data["spans"].append((deb_def, fin_def, plaque))

    # B. Intégration de SuiviJournalier :
    #    - EXCLUSIF pour aujourd'hui (j == jour_courant)
    #    - Repli pour jours passés non encore archivés dans HistoriqueJournalier
    for s in suivis:
        j = s.date_jour
        is_today = (j == jour_courant)
        plaque = s.vehicule.plaque if s.vehicule else None
        tcj_suivi = int(s.tcj_s or 0)
        ttj_suivi = int(s.ttj_s or 0)

        trips_par_gardien_suivi: dict[str, list[dict]] = {}
        if s.trajets:
            for t in s.trajets:
                if t.statut_validation == StatutValidationTrajet.REJETE:
                    continue
                if t.distance_km is not None and t.distance_km < 0.3:
                    continue

                deb = t.heure_debut
                if deb is None:
                    continue
                fin = t.heure_fin
                if fin is None:
                    fin_eff = maintenant if is_today else (deb + timedelta(minutes=15))
                else:
                    fin_eff = fin
                duree = int((fin_eff - deb).total_seconds()) if fin_eff >= deb else 0

                # Chauffeur spécifique du trajet :
                # 1. Badge sur le trajet
                t_gard = resoudre_gardien(t.conducteur_badge_id, t.conducteur_badge)
                # 2. Si pas de badge et que le camion a un titulaire officiel
                if not t_gard and s.vehicule and s.vehicule.conducteur_actuel_id:
                    t_gard = resoudre_gardien(s.vehicule.conducteur_actuel_id, None)
                # 3. Fallback sur le chauffeur du suivi
                if not t_gard:
                    t_gard = resoudre_gardien(s.conducteur_id, None)
                if not t_gard:
                    continue

                trips_par_gardien_suivi.setdefault(t_gard.id, []).append({
                    "gardien": t_gard,
                    "debut": deb,
                    "fin": fin_eff,
                    "duree": duree,
                    "plaque": plaque
                })

        if trips_par_gardien_suivi:
            nb_conducteurs = len(trips_par_gardien_suivi)
            if nb_conducteurs == 1:
                tgid, items = list(trips_par_gardien_suivi.items())[0]
                t_gard = items[0]["gardien"]
                if not conducteur_id_filtre or tgid == conducteur_id_filtre:
                    if is_today or (tgid, j) not in jours_archives_vus:
                        c_data = data_chauffeurs.setdefault(tgid, {"gardien": t_gard, "jours": {}, "spans": []})
                        for it in items:
                            c_data["spans"].append((it["debut"], it["fin"], plaque))
                        tcj_val = tcj_suivi if tcj_suivi > 0 else sum(it["duree"] for it in items)
                        ttj_val = max(ttj_suivi, tcj_val)
                        if j not in c_data["jours"]:
                            c_data["jours"][j] = {
                                "tcj_s": tcj_val, "ttj_s": ttj_val,
                                "vehicules": {plaque} if plaque else set(),
                                "en_cours": is_today
                            }
                        else:
                            c_data["jours"][j]["tcj_s"] += tcj_val
                            c_data["jours"][j]["ttj_s"] = max(c_data["jours"][j]["ttj_s"] + ttj_val, c_data["jours"][j]["tcj_s"])
                            if plaque:
                                c_data["jours"][j]["vehicules"].add(plaque)
                            if is_today:
                                c_data["jours"][j]["en_cours"] = True
            else:
                for tgid, items in trips_par_gardien_suivi.items():
                    if conducteur_id_filtre and tgid != conducteur_id_filtre:
                        continue
                    if not is_today and (tgid, j) in jours_archives_vus:
                        continue
                    t_gard = items[0]["gardien"]
                    c_data = data_chauffeurs.setdefault(tgid, {"gardien": t_gard, "jours": {}, "spans": []})
                    for it in items:
                        c_data["spans"].append((it["debut"], it["fin"], plaque))
                    tcj_c = sum(it["duree"] for it in items)
                    deb_min = min(it["debut"] for it in items)
                    fin_max = max(it["fin"] for it in items)
                    amp = int((fin_max - deb_min).total_seconds())
                    ttj_c = max(amp, tcj_c)
                    if j not in c_data["jours"]:
                        c_data["jours"][j] = {
                            "tcj_s": tcj_c, "ttj_s": ttj_c,
                            "vehicules": {plaque} if plaque else set(),
                            "en_cours": is_today
                        }
                    else:
                        c_data["jours"][j]["tcj_s"] += tcj_c
                        c_data["jours"][j]["ttj_s"] = max(c_data["jours"][j]["ttj_s"] + ttj_c, c_data["jours"][j]["tcj_s"])
                        if plaque:
                            c_data["jours"][j]["vehicules"].add(plaque)
                        if is_today:
                            c_data["jours"][j]["en_cours"] = True
        else:
            gardien = None
            if s.vehicule and s.vehicule.conducteur_actuel_id:
                gardien = resoudre_gardien(s.vehicule.conducteur_actuel_id, None)
            if not gardien:
                gardien = resoudre_gardien(s.conducteur_id, None)
            if gardien:
                gid = gardien.id
                if not conducteur_id_filtre or gid == conducteur_id_filtre:
                    if is_today or (gid, j) not in jours_archives_vus:
                        c_data = data_chauffeurs.setdefault(gid, {"gardien": gardien, "jours": {}, "spans": []})
                        tcj_val = tcj_suivi
                        ttj_val = max(ttj_suivi, tcj_val)

                        if j not in c_data["jours"]:
                            c_data["jours"][j] = {
                                "tcj_s": tcj_val, "ttj_s": ttj_val,
                                "vehicules": {plaque} if plaque else set(),
                                "en_cours": is_today
                            }
                        else:
                            c_data["jours"][j]["tcj_s"] += tcj_val
                            c_data["jours"][j]["ttj_s"] = max(c_data["jours"][j]["ttj_s"] + ttj_val, c_data["jours"][j]["tcj_s"])
                            if plaque:
                                c_data["jours"][j]["vehicules"].add(plaque)
                            if is_today:
                                c_data["jours"][j]["en_cours"] = True

                        if tcj_val > 0:
                            h_dep = s.heure_depart
                            if h_dep is not None:
                                fin_eff = h_dep + timedelta(seconds=max(1800, ttj_val or tcj_val))
                                c_data["spans"].append((h_dep, fin_eff, plaque))
                            else:
                                deb_def = datetime.combine(j, datetime.min.time()).replace(hour=6)
                                fin_def = deb_def + timedelta(seconds=max(1800, ttj_val or tcj_val))
                                c_data["spans"].append((deb_def, fin_def, plaque))

    # 4. Calcul de l'algorithme TCH (détection du reset 24h et cumul hebdomadaire)
    resultats_lignes = []

    # Génération de la liste des dates de la fenêtre demandée
    nb_jours = (fin_fenetre - debut_fenetre).days + 1
    dates_fenetre = [debut_fenetre + timedelta(days=i) for i in range(nb_jours)]
    dates_str = [d.isoformat() for d in dates_fenetre]

    for gid, c_data in data_chauffeurs.items():
        cond = c_data["gardien"]
        spans = c_data["spans"]
        jours_dict = c_data["jours"]

        # Fusion des intervalles qui se chevauchent
        spans.sort(key=lambda x: x[0])
        merged_spans: list[list[datetime]] = []
        for deb, fin, _ in spans:
            if fin < deb:
                fin = deb
            if not merged_spans:
                merged_spans.append([deb, fin])
            else:
                if deb <= merged_spans[-1][1]:
                    merged_spans[-1][1] = max(merged_spans[-1][1], fin)
                else:
                    merged_spans.append([deb, fin])

        # Recherche du dernier reset :
        # Période continue de repos >= 24h (86 400 s) sans aucun trajet valide.
        # Règle réglementaire TCH : un cycle hebdomadaire court sur 6 périodes de conduite max
        # (plafonné à 7 jours glissants).
        dernier_reset_debut: datetime | None = None
        date_dernier_reset_str: str | None = None
        tch_cumul_s = 0

        # Borne réglementaire hebdomadaire (au plus 6 jours avant aujourd'hui)
        borne_hebdo_date = jour_courant - timedelta(days=6)

        if not merged_spans:
            dernier_reset_debut = None
            date_dernier_reset_str = None
            tch_cumul_s = 0
        else:
            dernier_fin = merged_spans[-1][1]
            repos_depuis_fin = (maintenant - dernier_fin).total_seconds()

            if repos_depuis_fin >= SEUIL_RESET_REPOS_S:
                # Le chauffeur est au repos depuis plus de 24h jusqu'à maintenant :
                # Son TCH actif est remis à 0.
                dernier_reset_debut = None
                date_dernier_reset_str = dernier_fin.isoformat()
                tch_cumul_s = 0
            else:
                # Recherche de la dernière coupure >= 24h en remontant la chaîne
                dernier_reset_debut = merged_spans[0][0]
                date_dernier_reset_str = merged_spans[0][0].isoformat()

                for i in range(len(merged_spans) - 1, 0, -1):
                    gap = (merged_spans[i][0] - merged_spans[i - 1][1]).total_seconds()
                    if gap >= SEUIL_RESET_REPOS_S:
                        dernier_reset_debut = merged_spans[i][0]
                        date_dernier_reset_str = merged_spans[i][0].isoformat()
                        break

                # Borne de début de calcul : la plus récente entre le dernier reset 24h et le début du cycle hebdo (J-6)
                date_effective_debut = max(dernier_reset_debut.date(), borne_hebdo_date)

                # Somme des TCJ pour les journées à partir de la borne effective
                for j_date, j_info in jours_dict.items():
                    if j_date >= date_effective_debut and j_date <= jour_courant:
                        tch_cumul_s += int(j_info["tcj_s"] or 0)

        tch_restant_s = SEUIL_TCH_MAX_S - tch_cumul_s

        if tch_cumul_s >= SEUIL_TCH_MAX_S:
            alerte_statut = "LIMITE_ATTEINTE"
        elif tch_cumul_s >= SEUIL_TCH_ALERTE_S:
            alerte_statut = "PROCHE_LIMITE"
        else:
            alerte_statut = "NORMAL"

        # Construction de l'historique jour par jour pour la grille
        hist_grid = {}
        for d in dates_fenetre:
            d_str = d.isoformat()
            if d in jours_dict:
                j_info = jours_dict[d]
                inclus = (dernier_reset_debut is not None and d >= max(dernier_reset_debut.date(), borne_hebdo_date) and d <= jour_courant)
                hist_grid[d_str] = {
                    "tcj_s": j_info["tcj_s"],
                    "ttj_s": j_info["ttj_s"],
                    "vehicules": sorted(list(j_info["vehicules"])),
                    "inclus_dans_tch": inclus,
                    "en_cours": j_info.get("en_cours", False),
                }
            else:
                hist_grid[d_str] = {
                    "tcj_s": 0,
                    "ttj_s": 0,
                    "vehicules": [],
                    "inclus_dans_tch": False,
                    "en_cours": False,
                }

        # Véhicules actuellement actifs / récents
        vehicules_actifs = sorted(list(
            jours_dict.get(jour_courant, {}).get("vehicules", set())
        ))

        resultats_lignes.append({
            "conducteur_id": cond.id,
            "nom_prenom": cond.nom_prenom,
            "prenom_usuel": cond.prenom_usuel or "",
            "matricule": cond.matricule or "—",
            "telephone": cond.telephone,
            "statut": cond.statut.value if cond.statut else "ACTIF",
            "vehicules_actifs": vehicules_actifs,
            "tch_cumul_s": tch_cumul_s,
            "tch_restant_s": tch_restant_s,
            "date_dernier_reset": date_dernier_reset_str,
            "alerte_statut": alerte_statut,
            "historique": hist_grid,
        })

    # Tri par défaut : statut alerte (critique d'abord), puis TCH décroissant, puis nom
    ordre_alerte = {"LIMITE_ATTEINTE": 0, "PROCHE_LIMITE": 1, "NORMAL": 2}
    resultats_lignes.sort(key=lambda x: (
        ordre_alerte.get(x["alerte_statut"], 2),
        -x["tch_cumul_s"],
        (x["nom_prenom"] or "").upper()
    ))

    # Calcul des statistiques globales
    total_chauffeurs = len(resultats_lignes)
    en_conduite = sum(1 for l in resultats_lignes if l["historique"].get(jour_courant.isoformat(), {}).get("tcj_s", 0) > 0)
    proches_limite = sum(1 for l in resultats_lignes if l["alerte_statut"] == "PROCHE_LIMITE")
    limite_atteinte = sum(1 for l in resultats_lignes if l["alerte_statut"] == "LIMITE_ATTEINTE")
    cumuls_actifs = [l["tch_cumul_s"] for l in resultats_lignes if l["tch_cumul_s"] > 0]
    tch_moyen_s = int(sum(cumuls_actifs) / len(cumuls_actifs)) if cumuls_actifs else 0

    return {
        "du": debut_fenetre.isoformat(),
        "au": fin_fenetre.isoformat(),
        "dates": dates_str,
        "seuils": {
            "seuil_alerte_s": SEUIL_TCH_ALERTE_S,
            "seuil_max_s": SEUIL_TCH_MAX_S,
            "seuil_reset_repos_s": SEUIL_RESET_REPOS_S,
        },
        "stats": {
            "total_chauffeurs": total_chauffeurs,
            "en_conduite_aujourdhui": en_conduite,
            "proche_limite": proches_limite,
            "limite_atteinte": limite_atteinte,
            "tch_moyen_s": tch_moyen_s,
        },
        "lignes": resultats_lignes,
    }


def calculer_tch_seul_conducteur(db: Session, conducteur_id: str,
                                 maintenant: datetime | None = None) -> dict:
    """Calcul rapide du TCH pour un chauffeur donné (utilisé par le moteur d'alertes)."""
    maintenant = maintenant or now_local()
    jour = jour_attribution(maintenant)
    res = extraire_donnees_chauffeurs(db, jour - timedelta(days=6), jour,
                                     maintenant=maintenant,
                                     conducteur_id_filtre=conducteur_id)
    lignes = res.get("lignes") or []
    if lignes:
        return lignes[0]
    return {
        "conducteur_id": conducteur_id,
        "tch_cumul_s": 0,
        "tch_restant_s": SEUIL_TCH_MAX_S,
        "alerte_statut": "NORMAL",
        "date_dernier_reset": None,
    }


@router.get("")
def synthese_temps_conduite(du: str | None = None, au: str | None = None,
                            q: str | None = None, statut: str | None = None,
                            alerte: str | None = None,
                            db: Session = Depends(get_db),
                            _=Depends(require_roles(*TOUS))):
    """Vue principale — matrice des temps de conduite par chauffeur (TCH + TCJ/TTJ quotidiens)."""
    d_du, d_au = _determiner_periode(du, au)
    maintenant = now_local()
    res = extraire_donnees_chauffeurs(db, d_du, d_au, maintenant=maintenant)

    # Filtrage côté requête
    lignes = res["lignes"]
    if q and q.strip():
        motif = q.strip().lower()
        lignes = [
            l for l in lignes
            if motif in (l["prenom_usuel"] or "").lower()
            or motif in (l["nom_prenom"] or "").lower()
            or motif in (l["matricule"] or "").lower()
            or any(motif in v.lower() for v in l["vehicules_actifs"])
        ]
    if statut:
        lignes = [l for l in lignes if l["statut"] == statut]
    if alerte:
        lignes = [l for l in lignes if l["alerte_statut"] == alerte]

    res["lignes"] = lignes
    return res


@router.get("/conducteur/{cid}")
def detail_conducteur_tch(cid: str, du: str | None = None, au: str | None = None,
                          db: Session = Depends(get_db),
                          _=Depends(require_roles(*TOUS))):
    """Détail du TCH d'un chauffeur donné."""
    d_du, d_au = _determiner_periode(du, au)
    maintenant = now_local()
    res = extraire_donnees_chauffeurs(db, d_du, d_au, maintenant=maintenant,
                                     conducteur_id_filtre=cid)
    lignes = res.get("lignes") or []
    if not lignes:
        raise HTTPException(404, "Conducteur introuvable")
    return {"du": d_du.isoformat(), "au": d_au.isoformat(),
            "dates": res["dates"], "seuils": res["seuils"],
            "conducteur": lignes[0]}


# ============================== EXPORTS ==============================
BLEU_HEADER = "1F4E79"
BLEU_SUB = "2E75B6"
GRIS_LIGNE = "F9FBFD"
ORANGE_WARN = "FFF2CC"
ROUGE_CRITIQUE = "FCE4D6"
VERT_OK = "E2EFDA"


@router.get("/export.xlsx")
def export_temps_conduite_xlsx(du: str | None = None, au: str | None = None,
                               q: str | None = None, statut: str | None = None,
                               alerte: str | None = None,
                               db: Session = Depends(get_db),
                               _=Depends(require_roles(*TOUS))):
    """Export Excel de la matrice Temps de Conduite (TCH + TCJ/TTJ par jour)."""
    d_du, d_au = _determiner_periode(du, au)
    res = extraire_donnees_chauffeurs(db, d_du, d_au, maintenant=now_local())
    lignes = res["lignes"]

    if q and q.strip():
        motif = q.strip().lower()
        lignes = [
            l for l in lignes
            if motif in (l["prenom_usuel"] or "").lower()
            or motif in (l["nom_prenom"] or "").lower()
            or motif in (l["matricule"] or "").lower()
        ]
    if statut:
        lignes = [l for l in lignes if l["statut"] == statut]
    if alerte:
        lignes = [l for l in lignes if l["alerte_statut"] == alerte]

    dates_str = res["dates"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Temps de conduite (TCH)"

    # Ligne 1 : Titre
    nb_cols = 5 + len(dates_str) * 2
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=nb_cols)
    titre_cell = ws.cell(row=1, column=1,
                         value=f"Temps de Conduite Hebdomadaire (TCH) — Du {d_du:%d/%m/%Y} au {d_au:%d/%m/%Y} — Exporté le {datetime.now():%d/%m/%Y %H:%M}")
    titre_cell.font = Font(bold=True, size=13, color="FFFFFF")
    titre_cell.alignment = Alignment(horizontal="center", vertical="center")
    for col in range(1, nb_cols + 1):
        ws.cell(row=1, column=col).fill = PatternFill("solid", fgColor=BLEU_HEADER)
    ws.row_dimensions[1].height = 28

    # Ligne 2 : Super-headers
    headers_fixes = ["Chauffeur", "Matricule", "Statut", "Cumul TCH", "TCH restant"]
    for idx, h in enumerate(headers_fixes, start=1):
        ws.merge_cells(start_row=2, start_column=idx, end_row=3, end_column=idx)
        c = ws.cell(row=2, column=idx, value=h)
        c.font = Font(bold=True, size=11, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLEU_HEADER)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    col_idx = 6
    for d_str in dates_str:
        d_obj = date.fromisoformat(d_str)
        ws.merge_cells(start_row=2, start_column=col_idx, end_row=2, end_column=col_idx + 1)
        c = ws.cell(row=2, column=col_idx, value=d_obj.strftime("%d/%m/%Y"))
        c.font = Font(bold=True, size=10, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLEU_SUB)
        c.alignment = Alignment(horizontal="center", vertical="center")

        c_tcj = ws.cell(row=3, column=col_idx, value="TCJ")
        c_tcj.font = Font(bold=True, size=9)
        c_tcj.fill = PatternFill("solid", fgColor="E9EEF4")
        c_tcj.alignment = Alignment(horizontal="center", vertical="center")

        c_ttj = ws.cell(row=3, column=col_idx + 1, value="TTJ")
        c_ttj.font = Font(bold=True, size=9)
        c_ttj.fill = PatternFill("solid", fgColor="E9EEF4")
        c_ttj.alignment = Alignment(horizontal="center", vertical="center")

        col_idx += 2

    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 18

    # Lignes de données
    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3")
    )

    row_num = 4
    for l in lignes:
        # Col 1: Chauffeur
        nom_complet = f"{l['prenom_usuel']} ({l['nom_prenom']})" if l['prenom_usuel'] else l['nom_prenom']
        c_nom = ws.cell(row=row_num, column=1, value=nom_complet)
        c_nom.font = Font(bold=True)
        c_nom.alignment = Alignment(horizontal="left", vertical="center")

        # Col 2: Matricule
        c_mat = ws.cell(row=row_num, column=2, value=l['matricule'])
        c_mat.alignment = Alignment(horizontal="center", vertical="center")

        # Col 3: Statut
        c_stat = ws.cell(row=row_num, column=3, value=l['statut'])
        c_stat.alignment = Alignment(horizontal="center", vertical="center")

        # Col 4: Cumul TCH
        c_cumul = ws.cell(row=row_num, column=4, value=fmt_hms(l['tch_cumul_s']) or "00:00")
        c_cumul.font = Font(bold=True)
        c_cumul.alignment = Alignment(horizontal="center", vertical="center")
        if l['alerte_statut'] == "LIMITE_ATTEINTE":
            c_cumul.fill = PatternFill("solid", fgColor=ROUGE_CRITIQUE)
            c_cumul.font = Font(bold=True, color="9C0006")
        elif l['alerte_statut'] == "PROCHE_LIMITE":
            c_cumul.fill = PatternFill("solid", fgColor=ORANGE_WARN)
            c_cumul.font = Font(bold=True, color="9C6500")
        else:
            c_cumul.fill = PatternFill("solid", fgColor=VERT_OK)

        # Col 5: TCH restant
        c_rest = ws.cell(row=row_num, column=5, value=fmt_hms(l['tch_restant_s']) or "56:00")
        c_rest.alignment = Alignment(horizontal="center", vertical="center")
        if l['tch_restant_s'] <= 0:
            c_rest.font = Font(bold=True, color="9C0006")
        elif l['tch_restant_s'] <= 36000:  # <= 10h
            c_rest.font = Font(bold=True, color="9C6500")

        # Cols 6+: TCJ / TTJ
        c_col = 6
        for d_str in dates_str:
            h_data = l["historique"].get(d_str, {})
            tcj = h_data.get("tcj_s", 0)
            ttj = h_data.get("ttj_s", 0)

            c_tcj = ws.cell(row=row_num, column=c_col, value=fmt_hms(tcj) if tcj > 0 else "—")
            c_tcj.alignment = Alignment(horizontal="center", vertical="center")

            c_ttj = ws.cell(row=row_num, column=c_col + 1, value=fmt_hms(ttj) if ttj > 0 else "—")
            c_ttj.alignment = Alignment(horizontal="center", vertical="center")

            if h_data.get("en_cours"):
                c_tcj.fill = PatternFill("solid", fgColor="EBF1F5")
                c_ttj.fill = PatternFill("solid", fgColor="EBF1F5")

            c_col += 2

        # Bordures
        for col_idx_b in range(1, nb_cols + 1):
            ws.cell(row=row_num, column=col_idx_b).border = thin_border

        ws.row_dimensions[row_num].height = 20
        row_num += 1

    # Largeurs de colonnes
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    for c_i in range(6, nb_cols + 1):
        ws.column_dimensions[get_column_letter(c_i)].width = 9

    ws.freeze_panes = "F4"

    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=Temps_Conduite_{d_du}_{d_au}.xlsx"}
    )


@router.get("/export.pdf")
def export_temps_conduite_pdf(du: str | None = None, au: str | None = None,
                              q: str | None = None, statut: str | None = None,
                              alerte: str | None = None,
                              db: Session = Depends(get_db),
                              _=Depends(require_roles(*TOUS))):
    """Export PDF de la matrice Temps de Conduite."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    d_du, d_au = _determiner_periode(du, au)
    res = extraire_donnees_chauffeurs(db, d_du, d_au, maintenant=now_local())
    lignes = res["lignes"]

    if q and q.strip():
        motif = q.strip().lower()
        lignes = [
            l for l in lignes
            if motif in (l["prenom_usuel"] or "").lower()
            or motif in (l["nom_prenom"] or "").lower()
            or motif in (l["matricule"] or "").lower()
        ]
    if statut:
        lignes = [l for l in lignes if l["statut"] == statut]
    if alerte:
        lignes = [l for l in lignes if l["alerte_statut"] == alerte]

    dates_str = res["dates"]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=8, rightMargin=8, topMargin=10, bottomMargin=10
    )
    styles = getSampleStyleSheet()

    elements = []
    titre = f"<b>Temps de Conduite Hebdomadaire (TCH)</b> — Du {d_du:%d/%m/%Y} au {d_au:%d/%m/%Y}"
    elements.append(Paragraph(titre, styles["Title"]))
    elements.append(Spacer(1, 4))

    # Tableau PDF
    headers = ["Chauffeur", "Matricule", "Cumul TCH", "TCH restant"]
    for d_str in dates_str:
        d_obj = date.fromisoformat(d_str)
        headers.append(d_obj.strftime("%d/%m\nTCJ"))
        headers.append(d_obj.strftime("%d/%m\nTTJ"))

    table_data = [headers]
    for l in lignes[:60]:
        row = [
            l["prenom_usuel"] or l["nom_prenom"][:18],
            l["matricule"],
            fmt_hms(l["tch_cumul_s"]) or "00:00",
            fmt_hms(l["tch_restant_s"]) or "56:00",
        ]
        for d_str in dates_str:
            h_data = l["historique"].get(d_str, {})
            tcj = h_data.get("tcj_s", 0)
            ttj = h_data.get("ttj_s", 0)
            row.append(fmt_hms(tcj) if tcj > 0 else "—")
            row.append(fmt_hms(ttj) if ttj > 0 else "—")
        table_data.append(row)

    t = Table(table_data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("ALIGN", (0, 1), (0, -1), "LEFT"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D0D0D0")),
    ]))
    elements.append(t)
    doc.build(elements)

    return Response(
        buf.getvalue(),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=Temps_Conduite_{d_du}_{d_au}.pdf"}
    )
