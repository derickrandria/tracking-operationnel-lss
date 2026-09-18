# -*- coding: utf-8 -*-
"""Tests v1.53 — TTJ / TCC / TCJ et FUSION D'AFFICHAGE.

Clarification métier DÉFINITIVE du 18/09/2026 (arbitrages LSS) :
  §1 TTJ = temps écoulé depuis le 1ᵉʳ départ valide ; le seuil de 12 h est un
     SEUIL DE SIGNALEMENT (flag_ttj + rougeur), JAMAIS un plafond : le camion
     peut rouler au-delà, la durée réelle est conservée. Aucun seuil de 12 h
     n'est appliqué au TCC (logique différente).
  §2 trajet valide ⟺ distance ≥ 0,3 km. En dessous : conservé pour audit,
     jamais affiché, jamais dans les compteurs, jamais le départ.
  §3 fusion d'affichage = SÉQUENCES : interruption STRICTEMENT < 30 min →
     regroupement ; ≥ 30 min → clôture, pause affichée, séquence suivante.
     Tous les trajets valides restent conservés individuellement EN BASE.
     3 compteurs distincts ; `nb_trajets` n'a qu'UN sens (séquences affichées).
  §4 TCC = chrono de session ; toute interruption ≥ 30 min réinitialise à 0 ;
     le TCC reste à 0 jusqu'au prochain départ valide ; une interruption qui
     CONTIENT des mini-manœuvres s'analyse selon SA DURÉE (jamais par cumul) ;
     les mini-manœuvres ne deviennent jamais des trajets affichés.
  §5 journée close : `tcc_s` jamais écrasé par 0 ; « — » à l'interface ;
     « 0:00 » à l'export, accompagné d'une note de convention.

Exécution :
  cd backend
  DATABASE_URL="sqlite:////tmp/test_v153.db" python test_calcul_affichage_v153.py
"""
import io
import os
import sys
import shutil
from datetime import date, datetime, timedelta
from pathlib import Path

os.environ.setdefault("SIM_ENABLE", "0")

_URL = os.environ.get("DATABASE_URL", "")
if not _URL.startswith("sqlite:////tmp/"):
    print("⛔ Sécurité : DATABASE_URL doit être fourni et viser /tmp "
          f"(reçu : {_URL or 'non défini'})")
    sys.exit(1)

sys.path.insert(0, str(Path(__file__).resolve().parent))

from sqlalchemy import func, select                                   # noqa: E402

from app.chaines import Segment, construire_journee                   # noqa: E402
from app.database import SessionLocal                                 # noqa: E402
from app.engine import ensure_suivi, get_seuils, recalculer_temps     # noqa: E402
from app.main import APP_VERSION, migrer_schema                       # noqa: E402
from app.models import (HistoriqueJournalier, StatutSourceTrajet,     # noqa: E402
                        StatutValidationTrajet, SuiviJournalier, Trajet,
                        Vehicule)
from app.seed import seed_si_vide                                     # noqa: E402
from app.serializers import (compter_trajets_reels, fusionner_snapshot,  # noqa: E402
                             fusionner_trajets_affichage, s_historique,
                             s_suivi)

OK = KO = 0


def check(label, cond, detail=""):
    global OK, KO
    if cond:
        OK += 1
        print(f"  ✅ {label}")
    else:
        KO += 1
        print(f"  ❌ {label}   {detail}")


J = date(2026, 9, 17)


def dt(j, h, m=0, s=0):
    return datetime(j.year, j.month, j.day, h, m, s)


def seg(d0, d1, km, rejete=False):
    return Segment(debut=d0, fin=d1, distance_km=km, rejete=rejete)


# =============================================================================
print("=" * 74)
print("  LSS v1.53 — TTJ / TCC / TCJ & FUSION D'AFFICHAGE (clarification 18/09)")
print(f"  Base : {_URL}")
print("=" * 74)

# -----------------------------------------------------------------------------
print("\n[1] §3 — EXEMPLE OBLIGATOIRE : CALCUL")
segs = [seg(dt(J, 5, 0), dt(J, 5, 20), 0.299, rejete=True),   # manœuvre
        seg(dt(J, 5, 30), dt(J, 8, 0), 115.0),
        seg(dt(J, 8, 5), dt(J, 8, 15), 0.520),
        seg(dt(J, 8, 50), dt(J, 12, 0), 75.0)]
jr = construire_journee(segs, maintenant=dt(J, 12, 30), date_jour=J,
                        pause_min=1200, seuil_km=0.3, pause_affichee_min=1800)
check("3 trajets valides retenus (le 0,299 km n'est pas une ligne)",
      [l.distance_km for l in jr.lignes] == [115.0, 0.520, 75.0],
      f"{[l.distance_km for l in jr.lignes]}")
check("le 0,299 km n'alimente PAS le TCJ : 21 000 s = 9 000 + 600 + 11 400",
      jr.tcj_s == 21000, f"tcj={jr.tcj_s}")
check("les 5 minutes 08:00→08:05 sont DÉDUITES du TCJ (arrêt court déduit)",
      jr.tcj_s == (115.0 and 9000 + 600 + 11400),
      f"tcj={jr.tcj_s}")
check("TTJ = 05:30 → 12:00 (23 400 s)", jr.ttj_s == 23400, f"ttj={jr.ttj_s}")
check("arrêts = 2 400 s (5 min + 35 min) — la manœuvre n'ajoute rien",
      jr.ttj_s - jr.tcj_s == 2400, f"arrêts={jr.ttj_s - jr.tcj_s}")

# -----------------------------------------------------------------------------
print("\n[2] §3 — EXEMPLE OBLIGATOIRE : SÉQUENCES AFFICHÉES")
bruts = [{"heure_debut": l.debut.isoformat(), "heure_fin": l.fin.isoformat(),
          "distance_km": l.distance_km, "pause_apres_s": l.pause_apres_s,
          "segments": 1} for l in jr.lignes]
fus = fusionner_trajets_affichage(bruts, seuil_fusion_s=1800, seuil_pause_aff_s=1800)
check("2 séquences affichées", len(fus) == 2, f"{len(fus)}")
check("T1 = 05:30 → 08:15 (le trajet 08:05→08:15 est DANS le groupe)",
      fus[0]["heure_debut"][11:16] == "05:30" and fus[0]["heure_fin"][11:16] == "08:15",
      f"{fus[0]['heure_debut'][11:16]}→{fus[0]['heure_fin'][11:16]}")
check("T1 agrège la distance : 115 + 0,520 = 115,52 km",
      abs(fus[0]["distance_km"] - 115.52) < 1e-6, f"{fus[0]['distance_km']}")
check("pause affichée entre les séquences = 00:35 (2 100 s)",
      fus[0]["pause_apres_s"] == 2100, f"{fus[0]['pause_apres_s']}")
check("T2 commence à 08:50 (08:50 → 12:00)", fus[1]["heure_debut"][11:16] == "08:50")
check("3 trajets valides RÉELS en base", compter_trajets_reels(fus) == 3,
      f"{compter_trajets_reels(fus)}")
check("2 séquences + 1 trajet regroupé (3 = 2 + 1)",
      compter_trajets_reels(fus) - len(fus) == 1)
check("les 5 min (08:00→08:05) ne deviennent JAMAIS une pause affichée",
      all(f.get("pause_apres_s") != 300 for f in fus))

# -----------------------------------------------------------------------------
print("\n[3] §3 — CAS COMPLÉMENTAIRE : aucune interruption longue mais un trou ≥ 30 min")
jr2 = construire_journee([seg(dt(J, 5, 30), dt(J, 8, 0), 115.0),
                          seg(dt(J, 8, 31), dt(J, 10, 0), 60.0)],
                         maintenant=dt(J, 10, 30), date_jour=J,
                         pause_min=1200, seuil_km=0.3, pause_affichee_min=1800)
fus2 = fusionner_trajets_affichage(
    [{"heure_debut": l.debut.isoformat(), "heure_fin": l.fin.isoformat(),
      "distance_km": l.distance_km, "pause_apres_s": l.pause_apres_s, "segments": 1}
     for l in jr2.lignes], seuil_fusion_s=1800, seuil_pause_aff_s=1800)
check("T1 se termine à 08:00 (pas de regroupement)", fus2[0]["heure_fin"][11:16] == "08:00",
      fus2[0]["heure_fin"][11:16])
check("pause affichée = 00:31 (1 860 s)", fus2[0]["pause_apres_s"] == 1860,
      f"{fus2[0]['pause_apres_s']}")
check("T2 commence à 08:31", fus2[1]["heure_debut"][11:16] == "08:31")

# -----------------------------------------------------------------------------
print("\n[4] §3 — BORNE DE REGROUPEMENT : 29:59 fusionne, 30:00 clôt")
def sequences_pour_trou(gap_s):
    d2 = dt(J, 8, 0) + timedelta(seconds=gap_s)
    jj = construire_journee([seg(dt(J, 5, 30), dt(J, 8, 0), 115.0),
                             seg(d2, dt(J, 10, 0), 60.0)],
                            maintenant=dt(J, 10, 30), date_jour=J,
                            pause_min=1200, seuil_km=0.3, pause_affichee_min=1800)
    return fusionner_trajets_affichage(
        [{"heure_debut": l.debut.isoformat(), "heure_fin": l.fin.isoformat(),
          "distance_km": l.distance_km, "pause_apres_s": l.pause_apres_s,
          "segments": 1} for l in jj.lignes],
        seuil_fusion_s=1800, seuil_pause_aff_s=1800)

check("29:59 → UNE seule séquence (strictement < 30 min)", len(sequences_pour_trou(1799)) == 1)
check("30:00 → DEUX séquences (≥ 30 min clôt la séquence)", len(sequences_pour_trou(1800)) == 2)
check("30:01 → DEUX séquences", len(sequences_pour_trou(1801)) == 2)
check("dans la séquence fusionnée, la distance est AGRÉGÉE (115 + 60 = 175 km)",
      abs(sequences_pour_trou(1799)[0]["distance_km"] - 175.0) < 1e-6,
      f"{sequences_pour_trou(1799)[0]['distance_km']}")

# -----------------------------------------------------------------------------
print("\n[5] §4 — TCC : 29:59 poursuit la session, 30:00 la réinitialise")
seed_si_vide()
migrer_schema()
db = SessionLocal()
try:
    for modele in (Trajet, SuiviJournalier, HistoriqueJournalier):
        db.query(modele).delete()
    db.commit()
    v = db.scalars(select(Vehicule).order_by(Vehicule.plaque).limit(1)).first()

    def scenario(jour, plan, maintenant):
        """Crée un suivi + ses trajets et rend (tcc_s, heure_depart, ligne d'écran)."""
        s = ensure_suivi(db, v, jour)
        db.query(Trajet).filter(Trajet.suivi_id == s.id).delete()
        db.commit()
        for n, (d0, d1, km, rej) in enumerate(plan, start=1):
            db.add(Trajet(suivi_id=s.id, numero=n, heure_debut=d0, heure_fin=d1,
                          statut_source=StatutSourceTrajet.VALIDE,
                          source_plateforme="CAMTRACKPRO", distance_km=km,
                          statut_validation=(StatutValidationTrajet.REJETE if rej
                                             else StatutValidationTrajet.VALIDE)))
        db.commit()
        db.expire_all()
        s = db.get(SuiviJournalier, s.id)
        recalculer_temps(db, s, maintenant)
        # ⚠ `recalculer_temps` pose les compteurs en mémoire (le commit est à la
        # charge de l'appelant) : un `db.refresh()` ici RELIRAIT la base et
        # écraserait les valeurs calculées — on lit donc AVANT de committer.
        tcc, depart = int(s.tcc_s or 0), s.heure_depart
        db.commit()
        return tcc, depart

    J5 = date(2026, 9, 8)
    tcc_29, _ = scenario(J5, [(dt(J5, 6, 0), dt(J5, 7, 0), 40.0, False),
                              (dt(J5, 7, 29, 59), dt(J5, 8, 0), 30.0, False)],
                         dt(J5, 8, 0))
    check("interruption 29:59 → le chrono CONTINUE : TCC = 08:00 − 06:00 = 7 200 s",
          tcc_29 == 7200, f"tcc={tcc_29}")
    tcc_30, _ = scenario(J5, [(dt(J5, 6, 0), dt(J5, 7, 0), 40.0, False),
                              (dt(J5, 7, 30), dt(J5, 8, 0), 30.0, False)],
                         dt(J5, 8, 0))
    check("interruption 30:00 → TCC RÉINITIALISÉ : 08:00 − 07:30 = 1 800 s",
          tcc_30 == 1800, f"tcc={tcc_30}")

    # -------------------------------------------------------------------------
    print("\n[6] §4 — après réinitialisation : TCC à 0, puis repart au prochain départ valide")
    J6 = date(2026, 9, 9)
    s6 = ensure_suivi(db, v, J6)
    db.add(Trajet(suivi_id=s6.id, numero=1, heure_debut=dt(J6, 6, 0),
                  heure_fin=dt(J6, 7, 0), statut_source=StatutSourceTrajet.VALIDE,
                  source_plateforme="CAMTRACKPRO", distance_km=40.0,
                  statut_validation=StatutValidationTrajet.VALIDE))
    db.commit()
    db.expire_all()
    s6 = db.get(SuiviJournalier, s6.id)
    recalculer_temps(db, s6, dt(J6, 7, 40))      # 40 min après la fin : pause coupante
    tcc_6a = int(s6.tcc_s or 0)
    db.commit()
    check("pendant une interruption ≥ 30 min : TCC = 0 (le chrono est réinitialisé)",
          tcc_6a == 0, f"tcc={tcc_6a}")
    db.add(Trajet(suivi_id=s6.id, numero=2, heure_debut=dt(J6, 7, 45),
                  heure_fin=dt(J6, 7, 50), statut_source=StatutSourceTrajet.VALIDE,
                  source_plateforme="CAMTRACKPRO", distance_km=5.0,
                  statut_validation=StatutValidationTrajet.VALIDE))
    db.commit()
    db.expire_all()
    s6 = db.get(SuiviJournalier, s6.id)
    recalculer_temps(db, s6, dt(J6, 7, 50))
    tcc_6b = int(s6.tcc_s or 0)
    db.commit()
    check("au départ valide suivant : le chrono repart DE CE DÉPART "
          "(07:50 − 07:45 = 300 s, la pause n'est pas recomptée)",
          tcc_6b == 300, f"tcc={tcc_6b}")

    # -------------------------------------------------------------------------
    print("\n[7] §4 — une mini-manœuvre ne démarre PAS de session")
    J7 = date(2026, 9, 10)
    tcc7, dep7 = scenario(J7, [(dt(J7, 5, 0), dt(J7, 5, 20), 0.10, True),
                               (dt(J7, 6, 0), dt(J7, 7, 0), 40.0, False)],
                          dt(J7, 7, 10))
    check("le départ reste le 1ᵉʳ trajet VALIDE (06:00, pas 05:00)",
          dep7 == dt(J7, 6, 0), f"{dep7}")
    check("TCC = 07:10 − 06:00 = 4 200 s (la manœuvre ne crée pas de session)",
          tcc7 == 4200, f"tcc={tcc7}")
    sv7 = db.get(SuiviJournalier, ensure_suivi(db, v, J7).id)
    d7 = s_suivi(sv7, get_seuils(db))
    check("compteurs : 1 trajet valide réel, 1 séquence, la manœuvre n'apparaît pas",
          d7["nb_trajets_valides_reels"] == 1 and d7["nb_sequences_affichees"] == 1
          and d7["nb_trajets_fusionnes"] == 0 and d7["nb_trajets"] == 1,
          f"{d7['nb_trajets_valides_reels']}/{d7['nb_sequences_affichees']}")

    # -------------------------------------------------------------------------
    print("\n[8] §4 — interruption CONTENANT des mini-manœuvres : analysée par SA DURÉE")
    J8 = date(2026, 9, 11)
    # (A) une interruption de 32 min qui contient une manœuvre de 32 min
    tccA, _ = scenario(J8, [(dt(J8, 6, 0), dt(J8, 7, 0), 40.0, False),
                            (dt(J8, 7, 0), dt(J8, 7, 32), 0.10, True),
                            (dt(J8, 7, 32), dt(J8, 8, 0), 30.0, False)],
                       dt(J8, 8, 0))
    check("(A) interruption de 32 min ≥ 30 min → TCC réinitialisé (1 680 s)",
          tccA == 1680, f"tcc={tccA}")
    svA = db.get(SuiviJournalier, ensure_suivi(db, v, J8).id)
    dA = s_suivi(svA, get_seuils(db))
    check("(A) l'interruption s'AFFICHE en pause (32 min = 1 920 s)",
          dA["trajets"][0].get("pause_apres_s") == 1920,
          f"{dA['trajets'][0].get('pause_apres_s')}")
    check("(A) la manœuvre n'est jamais un trajet affiché (2 séquences)",
          dA["nb_sequences_affichees"] == 2 and dA["nb_trajets_valides_reels"] == 2,
          f"{dA['nb_sequences_affichees']}/{dA['nb_trajets_valides_reels']}")

    # (B) deux interruptions de 16 min : cumul 32 min mais AUCUNE ≥ 30 min
    J8b = date(2026, 9, 12)
    tccB, _ = scenario(J8b, [(dt(J8b, 6, 0), dt(J8b, 6, 30), 30.0, False),
                             (dt(J8b, 6, 30), dt(J8b, 6, 46), 0.10, True),
                             (dt(J8b, 6, 46), dt(J8b, 7, 2), 30.0, False),
                             (dt(J8b, 7, 2), dt(J8b, 7, 18), 0.10, True),
                             (dt(J8b, 7, 18), dt(J8b, 8, 0), 40.0, False)],
                        dt(J8b, 8, 10))
    check("(B) 2 × 16 min (32 min CUMULÉES) → le chrono N'EST PAS coupé : "
          "TCC = 08:10 − 06:00 = 7 800 s (arbitrage LSS du 18/09 : par interruption)",
          tccB == 7800, f"tcc={tccB}")
    svB = db.get(SuiviJournalier, ensure_suivi(db, v, J8b).id)
    dB = s_suivi(svB, get_seuils(db))
    check("(B) 3 trajets valides réels, 1 seule séquence affichée (rupture < 30 min)",
          dB["nb_trajets_valides_reels"] == 3 and dB["nb_sequences_affichees"] == 1
          and dB["nb_trajets_fusionnes"] == 2,
          f"{dB['nb_trajets_valides_reels']}/{dB['nb_sequences_affichees']}"
          f"/{dB['nb_trajets_fusionnes']}")

    # -------------------------------------------------------------------------
    print("\n[9] §1 — TTJ : le seuil de 12 h SIGNALE, il ne plafonne pas")
    J9 = date(2026, 9, 13)
    tcc9, dep9 = scenario(J9, [(dt(J9, 5, 0), dt(J9, 18, 30), 400.0, False)],
                          dt(J9, 18, 40))
    s9 = db.get(SuiviJournalier, ensure_suivi(db, v, J9).id)
    db.refresh(s9)                                  # les compteurs sont commités
    d9 = s_suivi(s9, get_seuils(db))
    check("TTJ stocké = 13 h 30 = 48 600 s, JAMAIS écrêté à 43 200",
          int(s9.ttj_s) == 48600, f"ttj={s9.ttj_s}")
    check("TTJ > 12 h → flag_ttj LEVÉ (signalement)",
          d9["flag_ttj"] is True and d9["ttj_s"] == 48600,
          f"flag={d9['flag_ttj']} ttj={d9['ttj_s']}")
    check("aucun seuil de 12 h sur le TCC : il vaut 13 h 40 = 49 200 s "
          "(> 43 200 — le TCC n'a pas de plafond de 12 h)",
          int(s9.tcc_s) == 49200, f"tcc={s9.tcc_s}")
    check("TCJ reste la conduite pure (13 h 30, la journée n'a qu'un trajet)",
          int(s9.tcj_s) == 48600, f"tcj={s9.tcj_s}")

    # -------------------------------------------------------------------------
    print("\n[10] §3/§5 — ARCHIVE : trajets bruts conservés + compteurs + TCC intact")
    J10 = date(2026, 9, 14)
    s10 = ensure_suivi(db, v, J10)
    db.query(Trajet).filter(Trajet.suivi_id == s10.id).delete()
    db.commit()
    for n, (d0, d1, km) in enumerate(((dt(J10, 21, 0), dt(J10, 22, 30), 60.0),
                                      (dt(J10, 22, 40), dt(J10, 23, 55), 30.0)), start=1):
        db.add(Trajet(suivi_id=s10.id, numero=n, heure_debut=d0, heure_fin=d1,
                      statut_source=StatutSourceTrajet.VALIDE,
                      source_plateforme="CAMTRACKPRO", distance_km=km,
                      statut_validation=StatutValidationTrajet.VALIDE))
    db.commit()
    from app.daily import recalculer_archives_journee
    recalculer_archives_journee(J10, db=db, rattraper_portail=False,
                                autoriser_reecriture=True, motif="test_v153")
    db.expire_all()
    h = db.scalar(select(HistoriqueJournalier).where(
        HistoriqueJournalier.date_jour == J10,
        HistoriqueJournalier.vehicule_id == v.id))
    snap = h.donnees or {}
    check("le snapshot conserve les trajets BRUTS : 2 trajets valides en base",
          len(snap.get("trajets") or []) == 2, f"{len(snap.get('trajets') or [])}")
    check("compteurs distincts écrits dans l'archive : 2 réels / 1 séquence / 1 regroupé",
          snap.get("nb_trajets_valides_reels") == 2
          and snap.get("nb_sequences_affichees") == 1
          and snap.get("nb_trajets_fusionnes") == 1
          and snap.get("nb_trajets") == 1,
          f"{snap.get('nb_trajets_valides_reels')}/{snap.get('nb_sequences_affichees')}"
          f"/{snap.get('nb_trajets_fusionnes')}/{snap.get('nb_trajets')}")
    check("`nb_trajets` = `nb_sequences_affichees` (un seul sens, jamais deux)",
          snap.get("nb_trajets") == snap.get("nb_sequences_affichees"))
    check("`tcc_s` de l'archive n'est PAS écrasé par 0",
          int(snap.get("tcc_s") or 0) == int(h.donnees.get("tcc_s") or 0)
          and int(snap.get("tcc_s") or 0) > 0, f"tcc={snap.get('tcc_s')}")
    relu = s_historique(h, detail=True, seuils=get_seuils(db))
    check("relecture : tcc_s conservé + drapeau d'affichage `tcc_masque`",
          int(relu["tcc_s"]) == int(snap["tcc_s"]) and relu.get("tcc_masque") is True,
          f"tcc={relu['tcc_s']} masque={relu.get('tcc_masque')}")
    check("relecture : les 4 compteurs remontent au client",
          relu["nb_trajets"] == 1 and relu["nb_sequences_affichees"] == 1
          and relu["nb_trajets_valides_reels"] == 2 and relu["nb_trajets_fusionnes"] == 1,
          f"{relu['nb_trajets']}/{relu['nb_sequences_affichees']}"
          f"/{relu['nb_trajets_valides_reels']}/{relu['nb_trajets_fusionnes']}")
    check("relecture : le snapshot STOCKÉ n'est jamais muté",
          (h.donnees or {}).get("nb_trajets_valides_reels") == 2)
    # archive ANCIENNE (sans les nouveaux champs) : les compteurs sont dérivés
    vieux = {"tcc_s": 5400, "tcj_s": 7200, "ttj_s": 10800, "nb_trajets": 1,
             "trajets": [{"heure_debut": "2026-09-14T21:00:00",
                          "heure_fin": "2026-09-14T23:55:00",
                          "distance_km": 90.0, "segments": 2}]}
    relu_vieux = fusionner_snapshot(vieux)
    check("archive ANTÉRIEURE à v1.53 : réel déduit de `segments` (2), sans migration",
          relu_vieux["nb_trajets_valides_reels"] == 2
          and relu_vieux["nb_trajets_fusionnes"] == 1
          and relu_vieux["nb_sequences_affichees"] == 1,
          f"{relu_vieux['nb_trajets_valides_reels']}"
          f"/{relu_vieux['nb_trajets_fusionnes']}")

    # -------------------------------------------------------------------------
    print("\n[11] §5 — EXPORTS : « 0:00 » + note de convention (donnée intacte)")
    from openpyxl import load_workbook
    from app.exporters import (NOTE_TCC_CONVENTION, export_excel,
                               export_suivi_excel)

    sv10 = db.get(SuiviJournalier, s10.id)
    ligne = s_suivi(sv10, get_seuils(db))
    ligne["tcc_masque"] = True                       # journée close (N1)
    xl = export_suivi_excel("14/09/2026", [ligne], detail=True, utilisateur="test")
    ws = load_workbook(io.BytesIO(xl)).active
    entetes = [(c.value or "") for c in ws[2]]
    idx = next(i for i, h in enumerate(entetes, start=1)
               if "tcc" in str(h).lower() or "conduite continue" in str(h).lower())
    valeurs_tcc = [ws.cell(row=r, column=idx).value for r in range(4, 4 + len([ligne]))]
    check("export Suivi : la cellule TCC d'une journée close vaut 0 (convention « 0:00 »)",
          all((v or 0) == 0 for v in valeurs_tcc), f"{valeurs_tcc}")
    texte = " | ".join(str(c.value) for row in ws.iter_rows() for c in row
                       if isinstance(c.value, str))
    check("export Suivi : la NOTE de convention est présente dans le fichier",
          "Convention d'affichage" in texte and "conservé en base" in texte)
    check("export Suivi : la note dit explicitement que la valeur est d'AFFICHAGE",
          NOTE_TCC_CONVENTION in texte)

    hx = export_excel("Historique 09/2026", ["Date", "TCC"], [["14/09/2026", "0:00"]],
                      note=NOTE_TCC_CONVENTION)
    wsh = load_workbook(io.BytesIO(hx)).active
    texteh = " | ".join(str(c.value) for row in wsh.iter_rows() for c in row
                        if isinstance(c.value, str))
    check("export Historique : note de convention également présente",
          "Convention d'affichage" in texteh)

    # --- v1.53 : les PDF (Suivi ET Historique) portent la même note ---------
    import pypdf
    from app.exporters import export_pdf, export_suivi_pdf

    pdf_suivi = export_suivi_pdf("14/09/2026", [ligne], detail=True, utilisateur="test")
    texte_pdf_suivi = " ".join(
        (pg.extract_text() or "") for pg in
        pypdf.PdfReader(io.BytesIO(pdf_suivi)).pages)
    check("export Suivi PDF : note de convention présente dans le texte du PDF",
          "Convention d'affichage" in texte_pdf_suivi)
    check("export Suivi PDF : TCC rendu « 0:00 » sur une journée close",
          "0:00" in texte_pdf_suivi, texte_pdf_suivi[:80])

    pdf_hist = export_pdf("Historique 09/2026", ["Date", "TCC"],
                          [["14/09/2026", "0:00"]], note=NOTE_TCC_CONVENTION)
    texte_pdf_hist = " ".join(
        (pg.extract_text() or "") for pg in
        pypdf.PdfReader(io.BytesIO(pdf_hist)).pages)
    check("export Historique PDF : note de convention présente",
          "Convention d'affichage" in texte_pdf_hist)
    check("la donnée TCC réelle reste disponible dans le snapshot exporté "
          "(l'export ne détruit rien)",
          int((h.donnees or {}).get("tcc_s") or 0) > 0)

    # -------------------------------------------------------------------------
    print("\n[12] §1 — BORNES TTJ : 12:00:00 INCLUSIF, plafond 24 h, valeur jamais écrasée")
    seuils12 = get_seuils(db)
    ttj_max12 = float(seuils12["SEUIL_TTJ_MAX"])
    check("le seuil TTJ paramétré vaut bien 12 h (43 200 s)", ttj_max12 == 43200,
          f"{ttj_max12}")

    J12a = date(2026, 9, 15)
    scenario(J12a, [(dt(J12a, 6, 0), dt(J12a, 18, 0), 100.0, False)], dt(J12a, 20, 0))
    s12a = db.get(SuiviJournalier, ensure_suivi(db, v, J12a).id)
    db.refresh(s12a)
    d12a = s_suivi(s12a, seuils12)
    check("TTJ = 12:00:00 PILE → SIGNALÉ (seuil INCLUSIF, clarification du 18/09)",
          int(d12a["ttj_s"]) == 43200 and d12a["flag_ttj"] is True,
          f"ttj={d12a['ttj_s']} flag={d12a['flag_ttj']}")

    J12b = date(2026, 9, 16)
    scenario(J12b, [(dt(J12b, 6, 0), dt(J12b, 17, 59, 59), 100.0, False)], dt(J12b, 20, 0))
    s12b = db.get(SuiviJournalier, ensure_suivi(db, v, J12b).id)
    db.refresh(s12b)
    d12b = s_suivi(s12b, seuils12)
    check("TTJ = 11:59:59 → NON signalé (juste sous le seuil)",
          int(d12b["ttj_s"]) == 43199 and d12b["flag_ttj"] is False,
          f"ttj={d12b['ttj_s']} flag={d12b['flag_ttj']}")

    # --- au-delà de 24 h : le bornage au jour clippe AVANT tout plafond
    _jr24 = construire_journee(
        [Segment(debut=dt(J, 5, 0), fin=dt(J + timedelta(days=1), 8, 0),
                 distance_km=900.0, rejete=False)],
        maintenant=dt(J, 23, 59, 59), date_jour=J, pause_min=1200, seuil_km=0.3)
    check("un trajet de 27 h est BORNÉ au jour J (05:00 → 23:59:59 = 68 399 s) "
          "AVANT tout plafond : aucune valeur > 24 h n'est produite",
          _jr24.ttj_s == 68399 and _jr24.ttj_s < 86400, f"ttj={_jr24.ttj_s}")
    _jr24b = construire_journee(
        [Segment(debut=dt(J, 0, 0), fin=dt(J, 23, 59, 59), distance_km=500.0,
                 rejete=False)],
        maintenant=dt(J, 23, 59, 59), date_jour=J, pause_min=1200, seuil_km=0.3)
    check("journée pleine 00:00 → 23:59:59 = 86 399 s : le maximum physiologique "
          "d'une journée civile reste sous la borne de 86 400 s",
          _jr24b.ttj_s == 86399)

    # --- le plafond technique ne peut donc écrêter qu'une donnée ANORMALE :
    #     on le prouve sur une archive fautive, et l'anomalie est DÉCLARÉE.
    h_faux = HistoriqueJournalier(
        date_jour=date(2026, 9, 5), annee=2026, mois=9, vehicule_id=v.id,
        donnees={"tcc_s": 5400, "tcj_s": 100000, "ttj_s": 120000,
                 "trajets": [], "nb_trajets": 0, "plaque": "ANOMALIE"},
        nb_infractions=0, nb_alertes=0)
    db.add(h_faux)
    db.commit()
    lu_faux = s_historique(h_faux, seuils=seuils12)
    check("valeur d'archive ANORMALE (TTJ 120 000 s = 33 h) : ramenée à la borne "
          "technique 24 h à la lecture", int(lu_faux["ttj_s"]) == 86400,
          f"ttj={lu_faux['ttj_s']}")
    check("…et l'anomalie est DÉCLARÉE au client (plus de correction silencieuse)",
          lu_faux.get("plafond_24h_applique") is True)
    check("l'archive fautive n'est PAS réécrite en base (lecture seule)",
          int((h_faux.donnees or {}).get("ttj_s") or 0) == 120000)
    check("TTJ réel jamais écrasé en base pour une journée NORMALE : "
          "la valeur stockée égale la valeur calculée",
          int(s12a.ttj_s) == 43200, f"stocké={s12a.ttj_s}")

    print("\n" + "=" * 74)
    print(f"  RÉSULTAT : {OK} OK / {KO} KO")
    print(f"  Version applicative : {APP_VERSION}")
    print("=" * 74)
finally:
    db.close()

for suffixe in ("", "-wal", "-shm"):
    chemin = _URL.replace("sqlite:///", "") + suffixe
    if os.path.exists(chemin):
        shutil.rmtree(chemin, ignore_errors=True) if os.path.isdir(chemin) else os.remove(chemin)

sys.exit(1 if KO else 0)
